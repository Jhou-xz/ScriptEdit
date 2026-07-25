import os
from io import BytesIO

from django.conf import settings
from docx import Document

from api.models import Block

CLIP_KIND_LABELS = {
    "muted": "Muted Background Clip",
    "clip": "Clip",
    "full": "Full Clip",
}


def _fmt_m_ss(seconds):
    if seconds is None:
        return None
    s = max(0, int(seconds))
    return f"{s // 60}:{s % 60:02d}"


def _clip_paren(block):
    label = CLIP_KIND_LABELS.get(block.clip_kind, "Clip")
    in_s = _fmt_m_ss(block.source_in_seconds)
    out_s = _fmt_m_ss(block.source_out_seconds)
    url = block.source_url or ""
    if block.clip_kind == "full":
        return f"(Full Clip {url})"
    if in_s and out_s:
        return f"({label} @ {in_s} - {out_s} {url})"
    if in_s:
        return f"({label} @ {in_s} {url})"
    return f"({label} {url})".replace("  ", " ").replace(" )", ")")


def _node_text(node):
    parts = []
    if node.get("type") == "text":
        parts.append(node.get("text", ""))
    for child in node.get("content", []) or []:
        parts.append(_node_text(child))
    return "".join(parts)


def _image_src(node):
    if node.get("type") == "image":
        return (node.get("attrs") or {}).get("src")
    for child in node.get("content", []) or []:
        found = _image_src(child)
        if found:
            return found
    return None


def _media_path(src):
    if not src or not src.startswith("/media/"):
        return None
    rel = src[len("/media/"):]
    path = os.path.join(settings.MEDIA_ROOT, rel)
    return path if os.path.exists(path) else None


def export_script_docx(script):
    doc = Document()
    doc.add_heading(script.title, level=0)

    vo_blocks = list(
        Block.objects.filter(track__script=script, track__is_script_track=True).order_by(
            "start_seconds"
        )
    )
    other_blocks = list(
        Block.objects.filter(track__script=script)
        .exclude(track__is_script_track=True)
        .order_by("start_seconds")
    )

    for vo in vo_blocks:
        if vo.title:
            doc.add_heading(vo.title, level=1)

        for node in (vo.content or {}).get("content", []) or []:
            ntype = node.get("type")
            text = _node_text(node).strip()
            if not text and ntype != "image":
                continue
            if ntype == "heading":
                level = node.get("attrs", {}).get("level", 2)
                doc.add_heading(text, level=min(level + 1, 4))
            elif ntype == "blockquote":
                doc.add_paragraph(f"“{text}”")
            elif ntype in ("bulletList", "orderedList"):
                style = "List Bullet" if ntype == "bulletList" else "List Number"
                for item in node.get("content", []) or []:
                    item_text = _node_text(item).strip()
                    if item_text:
                        doc.add_paragraph(item_text, style=style)
            elif text:
                doc.add_paragraph(text)

        if vo.editor_note:
            doc.add_paragraph(f"(Editor's note: {vo.editor_note})")

        related = [
            b
            for b in other_blocks
            if b.anchor_block_id == vo.id
            or (
                b.anchor_block_id is None
                and b.start_seconds < vo.start_seconds + vo.duration_seconds
                and b.start_seconds + b.duration_seconds > vo.start_seconds
            )
        ]
        for b in sorted(related, key=lambda x: x.start_seconds):
            if b.track.kind == "images":
                src = _image_src(b.content or {})
                path = _media_path(src)
                if path:
                    try:
                        doc.add_picture(path, width=None)
                    except Exception:
                        doc.add_paragraph(f"[image: {src}]")
            elif b.clip_kind:
                doc.add_paragraph(_clip_paren(b))
            elif b.source_url:
                note = f"{b.editor_note} " if b.editor_note else ""
                doc.add_paragraph(f"({note}{b.source_url})" if note else f"({b.source_url})")
            elif b.editor_note:
                doc.add_paragraph(f"({b.editor_note})")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
